from docx import Document

from tech6_cv_generator import FormBuilder


def test_employment_contact_dict(tmp_path):
    data = {
        "name": "Jane Q. Expert",
        "employment": [
            {
                "period": "2020",
                "organization": "Acme Consulting Ltd.",
                "title": "Lead Consultant",
                "contact": {"name": "Available upon request", "email": "", "phone": ""},
                "country": "Fantasia",
                "summary": "Did stuff",
            }
        ],
    }
    builder = FormBuilder(data)
    out = tmp_path / "test_employment.docx"
    builder.build(output_path=str(out))

    doc = Document(str(out))
    table = doc.tables[1]  # employment table, after personal info
    found_ref = any(
        "Available upon request" in cell.text for row in table.rows for cell in row.cells
    )
    found_highlight = any(
        "FFFF00" in run._r.xml
        for row in table.rows
        for cell in row.cells
        for para in cell.paragraphs
        for run in para.runs
    )
    assert found_ref
    assert found_highlight
