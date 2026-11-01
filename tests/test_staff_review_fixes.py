"""Tests for the bugs fixed in the 2.1.0 staff review.

Each test pins a specific pre-review behaviour so it cannot regress
silently: double builds, path traversal via expert names, mapper edge
cases, mapped-copy pollution, and CLI extraction failures.
"""

import json
from pathlib import Path

from docx import Document

from tech6_cv_generator import FormBuilder, map_to_canonical
from tech6_cv_generator.cli import main
from tech6_cv_generator.form_builder import is_placeholder, safe_filename_stem
from tech6_cv_generator.ingestor import extract_text_from_docx
from tech6_cv_generator.orchestrator import generate_one


def test_build_twice_produces_identical_documents(tmp_path):
    """A builder reused for a second build must not duplicate sections."""
    builder = FormBuilder({"name": "Jane Q. Expert"})
    first = builder.build(output_path=str(tmp_path / "a.docx"))
    second = builder.build(output_path=str(tmp_path / "b.docx"))

    def texts(path):
        doc = Document(path)
        return [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]

    assert texts(first) == texts(second)
    assert builder.report["output_path"] == second


def test_safe_filename_stem_blocks_path_traversal():
    assert safe_filename_stem("../../secret") == "secret"
    assert safe_filename_stem("a/b") == "a_b"
    assert safe_filename_stem("Jane Q. Expert") == "Jane_Q._Expert"
    assert safe_filename_stem("...") == "expert"


def test_build_with_traversal_name_stays_inside_output_dir(tmp_path):
    builder = FormBuilder({"name": "../../evil"})
    out = builder.build(output_dir=str(tmp_path))
    assert Path(out).parent == tmp_path
    assert Path(out).exists()


def test_is_placeholder_is_the_single_rule():
    assert is_placeholder("")
    assert is_placeholder("{Insert full name}")
    assert is_placeholder("Insert full name")
    assert not is_placeholder("Jane Q. Expert")


def test_language_without_proficiency_has_no_empty_parens():
    mapped = map_to_canonical(
        {"name": "Jane Q. Expert", "languages": [{"language": "English"}]}
    )
    assert mapped["languages"] == ["English"]


def test_nested_expert_profile_maps_instead_of_empty():
    mapped = map_to_canonical(
        {"wrapper": {"personal_info": {"full_name": "Jane Q. Expert"}, "education": []}}
    )
    assert mapped["name"] == "Jane Q. Expert"


def test_authorized_representative_is_a_name_not_a_date():
    mapped = map_to_canonical(
        {
            "expert_profile": {
                "personal_info": {"full_name": "Jane Q. Expert"},
                "certification_statement": {
                    "representative_name": "Acme Consulting Ltd.",
                    "representative_signature_date": "2024-01-01",
                },
            }
        }
    )
    assert mapped["authorized_representative"] == "Acme Consulting Ltd."


def test_generate_one_does_not_pollute_input_dir(tmp_path):
    json_dir = tmp_path / "inputs"
    json_dir.mkdir()
    kosovo = json_dir / "expert.json"
    kosovo.write_text(
        json.dumps({"expert_profile": {"personal_info": {"full_name": "Jane Q. Expert"}}}),
        encoding="utf-8",
    )
    entry = generate_one(kosovo, tmp_path / "outputs")
    assert entry["status"] == "ok"
    assert list(json_dir.iterdir()) == [kosovo]


def test_docx_extraction_keeps_document_order(tmp_path):
    doc_path = tmp_path / "ordered.docx"
    doc = Document()
    doc.add_paragraph("first paragraph")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "table row"
    doc.add_paragraph("last paragraph")
    doc.save(str(doc_path))

    text = extract_text_from_docx(doc_path)
    assert text.index("first paragraph") < text.index("table row") < text.index("last paragraph")


def test_cli_extract_corrupt_pdf_reports_error_not_traceback(tmp_path, capsys):
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"not a pdf at all")
    assert main(["extract", str(bad)]) == 1
    captured = capsys.readouterr()
    assert "Traceback" not in captured.err
    assert "bad.pdf" in captured.err
