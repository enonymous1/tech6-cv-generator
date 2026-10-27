from docx import Document

from tech6_cv_generator import FormBuilder


def test_personal_info_highlights(tmp_path):
    data = {
        "position_title": "",
        "name": "",
        "dob": "",
        "citizenship": "",
    }
    builder = FormBuilder(data)
    out = tmp_path / "test_personal.docx"
    builder.build(output_path=str(out))

    doc = Document(str(out))
    table = doc.tables[0]
    highlighted = any(
        "FFFF00" in run._r.xml
        for row in table.rows
        for cell in row.cells
        for para in cell.paragraphs
        for run in para.runs
    )
    assert highlighted
